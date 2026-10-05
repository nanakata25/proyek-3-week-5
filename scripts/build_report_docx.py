"""Create an editable Word report from the same content as build_report.py.

python scripts/build_report_docx.py --nim STUDENT_ID --name "Full Name"
Requires python-docx, beautifulsoup4, Pillow. No page is flattened to an image.
"""
from pathlib import Path
import argparse, json, re
from xml.sax.saxutils import escape
from bs4 import BeautifulSoup, NavigableString
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

root=Path(__file__).resolve().parents[1]
ap=argparse.ArgumentParser(); ap.add_argument('--nim',required=True); ap.add_argument('--name',required=True); args=ap.parse_args()
stem=args.nim+'_'+re.sub(r'[^\w-]+','_',args.name)+'_Laporan_Modul4'
filename=root/'output/pdf'/f'{stem}.pdf'
shots=root/'evidence/screenshots'
db=json.loads((root/'evidence/database-snapshot.json').read_text())
shop=json.loads((root/'evidence/shop-browser.json').read_text())
order=shop['after']['queries']['orders']['rows'][-1]
repo='https://github.com/nanakata25/proyek-3-week-5'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=Inches(.60);sec.left_margin=sec.right_margin=Inches(.60)
for name,size in [('Normal',9.5),('Title',25),('Heading 1',18),('Heading 2',11.5),('Caption',8)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(8);s.paragraph_format.line_spacing=1.10
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.styles['Caption'].font.italic=False
for style in doc.styles:
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
fig=0;story=[]

def rich(par,text):
    soup=BeautifulSoup(text.replace('<link ','<a ').replace('</link>','</a>'),'html.parser')
    def walk(node,bold=False,italic=False):
        if isinstance(node,NavigableString):
            r=par.add_run(str(node));r.bold=bold;r.italic=italic
        elif node.name=='br':par.add_run().add_break()
        elif node.name=='a' and node.get('href'):
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),par.part.relate_to(node['href'],RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'087E80');pr.append(color);r.append(pr)
            text=OxmlElement('w:t');text.text=node.get_text();r.append(text);link.append(r);par._p.append(link)
        else:
            for child in node.children:walk(child,bold or node.name=='b',italic or node.name=='i')
    for child in soup.contents:walk(child)

def p(s,style='BodyCustom'):
    mapping={'TitleCustom':'Title','HeadingCustom':'Heading 1','SubCustom':'Heading 2','CaptionCustom':'Caption'}
    par=doc.add_paragraph(style=mapping.get(style,'Normal'))
    if style in ('HeadingCustom','SubCustom'):
        s=re.sub(r'^([0-9]+)\.\s*',r'\1 ',s).replace(' - ',' ').replace(' / ',' ')
    rich(par,s)
    if style=='SmallCustom':
        for run in par.runs:run.font.size=Pt(8)
        par.paragraph_format.space_after=Pt(6)
    if style=='CodeCustom':
        for run in par.runs:run.font.name='Consolas';run.font.size=Pt(8)
    story.append(par)

def h(s):p(s,'SubCustom')
def page(title,kicker='MODUL 04 / LAPORAN PRAKTIKUM'):
    if story:doc.add_page_break()
    p(kicker,'SmallCustom');p(title,'HeadingCustom')

def table(headers,rows,widths=None,size=8.2):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [505/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Pt(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate([headers]+list(rows)):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'163846' if i==0 else ('F2F5F6' if i%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side,v in [('top',70),('bottom',70),('left',100),('right',100)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            par=cell.paragraphs[0];par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1.05
            run=par.add_run(str(value));run.font.size=Pt(size);run.bold=i==0
            if i==0:run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def pic(file,caption,maxh=235,width=505):
    global fig
    path=shots/file
    with Image.open(path)as im:w,h=im.size
    scale=min(width/w,maxh/h);fig+=1
    par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next=True;par.paragraph_format.space_after=Pt(4)
    par.add_run().add_picture(str(path),width=Pt(w*scale),height=Pt(h*scale))
    p(f'<b>Gambar {fig}.</b> '+caption,'CaptionCustom')

def code(s):p(escape(s).replace('\n','<br/>'),'CodeCustom')

# Execute only the shared report's declarative content, with Word equivalents.
source=(root/'scripts/build_report.py').read_text(encoding='utf-8')
content=source[source.index("p('PRAKTIKUM PEMROGRAMAN"):source.index('\ndef decorate')]
content=re.sub(r'story\.append\(Spacer\(1,\d+\)\);?', '', content)
exec(compile(content,'shared-report-content','exec'))
footer=sec.footer.paragraphs[0]
footer.paragraph_format.space_after=Pt(0)
r=footer.add_run('MODUL 04 | '+args.nim+' | '+args.name+'   ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Laporan Modul 4 Penyimpanan Data pada Web'
doc.core_properties.author=args.name
output=root/'output/docx'/f'{stem}.docx';output.parent.mkdir(parents=True,exist_ok=True)
doc.save(output);print(output)
